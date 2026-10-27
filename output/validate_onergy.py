#!/usr/bin/env python3
"""Offline validation of ONERGY's PPTX and LibreOffice-rendered PDF.
Run from any directory using the prepared minimal Python environment.
"""
import json
import re
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
import fitz
from pptx import Presentation

REPO=Path(__file__).resolve().parent.parent
PROJECT=REPO/'projects/onergy_europe_om_ppt169_20261027'
OUT=REPO/'output'
def norm(value): return re.sub(r'\s+','',value)
def walk(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape,'shapes'): yield from walk(shape.shapes)
def main():
    pptx=OUT/'onergy_europe_om_draft.pptx';pdf=OUT/'onergy_europe_om_draft.pdf'
    deck=Presentation(pptx);doc=fitz.open(pdf)
    assert len(deck.slides)==len(doc)==20
    with zipfile.ZipFile(pptx) as z:
        assert z.testzip() is None
        notes=[n for n in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)]
        assert len(notes)==20
    report={'pages':20,'aspect_ratio':'16:9','slide_checks':[],'missing_text':[],'unexpected_wrapping':[],'out_of_bounds':[]}
    for index,(slide,pg,svg) in enumerate(zip(deck.slides,doc,sorted((PROJECT/'svg_output').glob('*.svg'))),1):
        shapes=list(walk(slide.shapes));text=[sh for sh in shapes if sh.has_text_frame and sh.text.strip()]
        assert text
        for sh in text:
            assert sh.text_frame.word_wrap is False
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    assert run.font.name=='Noto Sans CJK SC'
        report['slide_checks'].append({'page':index,'editable_text_objects':len(text),'picture_objects':sum(sh.shape_type==13 for sh in shapes),'speaker_notes':slide.has_notes_slide})
        lines=[]
        for block in pg.get_text('dict')['blocks']:
            for ln in block.get('lines',[]):
                lines.append(norm(''.join(sp['text'] for sp in ln['spans'])))
                for sp in ln['spans']:
                    x0,y0,x1,y1=sp['bbox']
                    if x0<-.5 or y0<-.5 or x1>pg.rect.width+.5 or y1>pg.rect.height+.5:
                        report['out_of_bounds'].append({'page':index,'text':sp['text']})
        actual=norm(pg.get_text())
        for el in ET.parse(svg).getroot().iter('{http://www.w3.org/2000/svg}text'):
            t=norm(''.join(el.itertext()))
            if t not in actual: report['missing_text'].append({'page':index,'text':t})
            if not any(t in ln for ln in lines):report['unexpected_wrapping'].append({'page':index,'text':t})
    # Verify text can actually be edited and round-tripped, without touching deliverables.
    with tempfile.TemporaryDirectory(prefix='onergy-edit-') as tmp:
        first=next(sh for sh in walk(deck.slides[0].shapes) if sh.has_text_frame and sh.text.strip())
        original=first.text;first.text='编辑验证 / editable'
        target=Path(tmp)/'edited.pptx';deck.save(target)
        again=Presentation(target)
        assert any(sh.has_text_frame and sh.text=='编辑验证 / editable' for sh in walk(again.slides[0].shapes))
    report['edit_roundtrip']='passed'
    report['fonts_in_pdf']=sorted({font[3].split('+')[-1] for pg in doc for font in pg.get_fonts()})
    (OUT/'onergy_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps({k:v for k,v in report.items() if k!='slide_checks'},ensure_ascii=False,indent=2))
    assert not report['missing_text'] and not report['out_of_bounds'] and not report['unexpected_wrapping']
if __name__=='__main__':main()
